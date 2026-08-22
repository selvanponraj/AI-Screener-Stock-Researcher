"""Evaluate SSGR and CFO/EBITDA rules from Screener Excel exports."""

from __future__ import annotations

import argparse
import csv
import json
import math
from pathlib import Path
from statistics import mean
from typing import Any, Optional

from openpyxl import load_workbook


DEFAULT_COMPANY_DIR = Path("data") / "runs" / "20260625_updated_company_profiles" / "companies"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Evaluate SSGR and CFO/EBITDA rules from downloaded Screener Excel files."
    )
    parser.add_argument(
        "--excel-dir",
        type=Path,
        default=DEFAULT_COMPANY_DIR / "excel",
        help="Directory containing the downloaded Excel files and manifest.",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=None,
        help="Directory for Excel-rule CSV and JSON. Defaults to <excel-dir>/analysis.",
    )
    return parser.parse_args()


def number(value: Any) -> Optional[float]:
    if isinstance(value, bool) or value is None:
        return None
    try:
        parsed = float(value)
    except (TypeError, ValueError):
        return None
    return parsed if math.isfinite(parsed) else None


def data_row(sheet: Any, row: int) -> list[Optional[float]]:
    return [number(sheet.cell(row, column).value) for column in range(2, 12)]


def average(values: list[Optional[float]]) -> Optional[float]:
    return mean(values) if all(value is not None for value in values) else None


def safe_ratio(numerator: Optional[float], denominator: Optional[float]) -> Optional[float]:
    if numerator is None or denominator is None or denominator == 0:
        return None
    return numerator / denominator


def zero_blank(value: Optional[float]) -> float:
    return 0.0 if value is None else value


def operating_profit(data_sheet: Any) -> list[Optional[float]]:
    sales = data_row(data_sheet, 17)
    raw_material = data_row(data_sheet, 18)
    inventory_change = data_row(data_sheet, 19)
    power_fuel = data_row(data_sheet, 20)
    manufacturing = data_row(data_sheet, 21)
    employees = data_row(data_sheet, 22)
    selling = data_row(data_sheet, 23)
    other_expenses = data_row(data_sheet, 24)

    values: list[Optional[float]] = []
    for index in range(len(sales)):
        if sales[index] is None:
            values.append(None)
            continue
        # Matches Profit & Loss!5: Excel SUM ignores blank expense cells.
        expenses = (
            zero_blank(raw_material[index])
            - zero_blank(inventory_change[index])
            + zero_blank(power_fuel[index])
            + zero_blank(manufacturing[index])
            + zero_blank(employees[index])
            + zero_blank(selling[index])
            + zero_blank(other_expenses[index])
        )
        values.append(sales[index] - expenses)
    return values


def trailing_operating_profit(data_sheet: Any) -> Optional[float]:
    quarter_operating_profit = [
        number(data_sheet.cell(50, column).value) for column in range(8, 12)
    ]
    return sum(quarter_operating_profit) if all(value is not None for value in quarter_operating_profit) else None


def ssgr_values(data_sheet: Any) -> list[Optional[float]]:
    sales = data_row(data_sheet, 17)
    net_profit = data_row(data_sheet, 30)
    dividend = data_row(data_sheet, 31)
    depreciation = data_row(data_sheet, 26)
    net_block = data_row(data_sheet, 62)

    # Profit & Loss!31 stores NPM as a percentage, not a decimal fraction.
    npm = [
        safe_ratio(profit, sale) * 100
        if safe_ratio(profit, sale) is not None
        else None
        for profit, sale in zip(net_profit, sales)
    ]
    # Profit & Loss!18 is IF(net_profit > 0, dividend / net_profit, 0).
    # Excel treats blank dividend cells as zero in the division.
    dpr = [
        0.0
        if profit is None or profit <= 0
        else safe_ratio(zero_blank(dividend_value), profit)
        for dividend_value, profit in zip(dividend, net_profit)
    ]
    depreciation_rate = [
        safe_ratio(depreciation_value, block) * 100
        if safe_ratio(depreciation_value, block) is not None
        else None
        for depreciation_value, block in zip(depreciation, net_block)
    ]
    nfat = [
        safe_ratio(sales[index], average([net_block[index], net_block[index - 1]]))
        if index > 0
        else None
        for index in range(len(sales))
    ]

    values: list[Optional[float]] = []
    # Excel cells E27:K27 correspond to annual-data indexes 3 through 9.
    for index in range(3, 10):
        npm_average = average(npm[index - 2 : index + 1])
        dpr_average = average(dpr[index - 2 : index + 1])
        nfat_average = average(nfat[index - 2 : index + 1])
        depreciation_average = average(depreciation_rate[index - 2 : index + 1])
        if any(value is None for value in (npm_average, dpr_average, nfat_average, depreciation_average)):
            values.append(None)
            continue
        values.append(nfat_average * npm_average * (1 - dpr_average) - depreciation_average)
    return values


def analyze_workbook(path: Path) -> dict[str, Any]:
    # openpyxl does not calculate formulas. Read the source Data Sheet and recreate
    # the template formulas so results do not depend on cached Excel values.
    workbook = load_workbook(path, data_only=False, read_only=True)
    try:
        data_sheet = workbook["Data Sheet"]
        ssgr = ssgr_values(data_sheet)
        ssgr_3y_latest = [v for v in ssgr[-3:] if v is not None]
        ssgr_3y_avg = average(ssgr_3y_latest) if len(ssgr_3y_latest) == 3 else None

        sales = data_row(data_sheet, 17)
        sales_cagr_3y: float | None = None
        if len(sales) >= 4 and sales[-4] is not None and sales[-1] is not None and sales[-4] > 0 and sales[-1] > 0:
            sales_cagr_3y = ((sales[-1] / sales[-4]) ** (1 / 3) - 1) * 100

        ssgr_pass = (
            ssgr_3y_avg > 10.0 and sales_cagr_3y is not None and ssgr_3y_avg >= sales_cagr_3y
            if ssgr_3y_avg is not None and sales_cagr_3y is not None
            else False
        )
        ssgr_rule = "missing" if ssgr_3y_avg is None or sales_cagr_3y is None else ("pass" if ssgr_pass else "fail")

        cfo = data_row(data_sheet, 82)
        net_profit = data_row(data_sheet, 30)
        operating_profit_values = operating_profit(data_sheet)
        ebitda = operating_profit_values[1:] + [trailing_operating_profit(data_sheet)]

        cfo_5y = [v for v in cfo[-5:] if v is not None]
        ebitda_5y = [v for v in ebitda[-5:] if v is not None]
        pat_5y = [v for v in net_profit[-5:] if v is not None]

        cum_cfo_5y = sum(cfo_5y) if len(cfo_5y) == 5 else None
        cum_ebitda_5y = sum(ebitda_5y) if len(ebitda_5y) == 5 else None
        cum_pat_5y = sum(pat_5y) if len(pat_5y) == 5 else None

        cfo_ebitda_cum_ratio: float | None = None
        if cum_cfo_5y is not None and cum_ebitda_5y is not None and cum_ebitda_5y > 0:
            cfo_ebitda_cum_ratio = (cum_cfo_5y / cum_ebitda_5y) * 100

        cfo_pat_cum_ratio: float | None = None
        if cum_cfo_5y is not None and cum_pat_5y is not None and cum_pat_5y > 0:
            cfo_pat_cum_ratio = (cum_cfo_5y / cum_pat_5y) * 100

        cfo_pass = (
            cfo_ebitda_cum_ratio >= 65.0 and cfo_pat_cum_ratio >= 80.0
            if cfo_ebitda_cum_ratio is not None and cfo_pat_cum_ratio is not None
            else False
        )

        cfo_rule = (
            "missing"
            if cfo_ebitda_cum_ratio is None or cfo_pat_cum_ratio is None
            else ("pass" if cfo_pass else "fail")
        )

        return {
            "ssgr_values_e27_k27": ssgr,
            "ssgr_3y_avg": ssgr_3y_avg,
            "excel_sales_cagr_3y": sales_cagr_3y,
            "rule_12_ssgr": ssgr_rule,
            "cum_cfo_5y": cum_cfo_5y,
            "cum_ebitda_5y": cum_ebitda_5y,
            "cum_pat_5y": cum_pat_5y,
            "cfo_ebitda_cum_ratio": cfo_ebitda_cum_ratio,
            "cfo_pat_cum_ratio": cfo_pat_cum_ratio,
            "rule_13_cfo_ebitda": cfo_rule,
        }
    finally:
        workbook.close()


def main() -> int:
    args = parse_args()
    excel_dir = args.excel_dir.resolve()
    manifest_path = excel_dir / "manifest.json"
    output_dir = (args.output_dir or excel_dir / "analysis").resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    if not manifest_path.is_file():
        print(f"Notice: Excel manifest not found at {manifest_path}. Skipping Excel rule analysis.")
        return 0

    downloads = json.loads(manifest_path.read_text(encoding="utf-8"))
    results: list[dict[str, Any]] = []
    for download in downloads:
        filename = download.get("download_file")
        if not filename or not (excel_dir / filename).is_file():
            continue
        result = {
            "company_name": download["company_name"],
            "company_url": download["company_url"],
            "prior_rule_pass_count": int(download["prior_rule_pass_count"]),
            "excel_file": filename,
            **analyze_workbook(excel_dir / filename),
        }
        results.append(result)

    json_path = output_dir / "excel_rule_results.json"
    csv_path = output_dir / "excel_rule_results.csv"
    json_path.write_text(json.dumps(results, indent=2), encoding="utf-8")
    fieldnames = list(results[0]) if results else []
    with csv_path.open("w", newline="", encoding="utf-8") as csv_file:
        writer = csv.DictWriter(csv_file, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(results)

    print(f"Analyzed {len(results)} Excel exports.")
    print(f"SSGR passes: {sum(result['rule_12_ssgr'] == 'pass' for result in results)}.")
    print(f"CFO/EBITDA passes: {sum(result['rule_13_cfo_ebitda'] == 'pass' for result in results)}.")
    print(f"Saved CSV: {csv_path}")
    print(f"Saved JSON: {json_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
